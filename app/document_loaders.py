"""Loaders for non-image document formats: PDF, DOCX, XLSX.

PDF pages are rendered to images and pushed through the OCR pipeline.
DOCX / XLSX are read natively (they are text files, not scans), and their
paragraphs/sheets are converted into the same text-block / table structures
the OCR engine returns, so the API response shape stays uniform.
"""

import io
import logging
from typing import Optional

from PIL import Image

logger = logging.getLogger(__name__)

# Magic-byte signatures used to detect format regardless of file extension
_PDF_MAGIC = b"%PDF-"
_ZIP_MAGIC = b"PK\x03\x04"  # docx/xlsx are both ZIP containers


def detect_document_type(contents: bytes, filename: str = "") -> str:
    """
    Detect the document type from file contents (falls back to extension).

    Returns one of: pdf, docx, xlsx, image, unsupported.
    """
    name = (filename or "").lower()

    if contents.startswith(_PDF_MAGIC):
        return "pdf"

    if contents.startswith(_ZIP_MAGIC):
        # Both .docx and .xlsx are ZIP; distinguish by their internal manifest
        try:
            import zipfile
            with zipfile.ZipFile(io.BytesIO(contents)) as zf:
                names = set(zf.namelist())
            if "word/document.xml" in names:
                return "docx"
            if any(n.startswith("xl/") for n in names):
                return "xlsx"
        except Exception:
            pass
        # Fall back to extension if the ZIP internals were unrecognized
        if name.endswith(".docx"):
            return "docx"
        if name.endswith(".xlsx"):
            return "xlsx"
        return "unsupported"

    if name.endswith(".pdf"):
        return "pdf"
    if name.endswith(".docx"):
        return "docx"
    if name.endswith(".xlsx"):
        return "xlsx"

    try:
        Image.open(io.BytesIO(contents)).verify()
        return "image"
    except Exception:
        return "unsupported"


def pdf_page_count(contents: bytes) -> int:
    """Return the number of pages in a PDF."""
    import fitz  # PyMuPDF
    with fitz.open(stream=contents, filetype="pdf") as doc:
        return doc.page_count


def pdf_pages_to_images(
    contents: bytes,
    dpi: int = 200,
    max_pages: Optional[int] = None,
) -> list[Image.Image]:
    """
    Render each PDF page to a PIL RGB image for OCR processing.

    Args:
        contents: Raw PDF bytes.
        dpi: Render resolution (200 dpi is a good OCR/size tradeoff).
        max_pages: Optional cap on number of pages processed.
    """
    import fitz  # PyMuPDF
    images = []
    with fitz.open(stream=contents, filetype="pdf") as doc:
        limit = doc.page_count if max_pages is None else min(doc.page_count, max_pages)
        for i in range(limit):
            page = doc[i]
            pix = page.get_pixmap(dpi=dpi)
            img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            images.append(img)
    return images


# ---------------------------------------------------------------------------
# Native Office loaders (DOCX / XLSX)
# ---------------------------------------------------------------------------

def _block(text: str) -> dict:
    """Build a text block dict compatible with ocr_engine.extract_text output."""
    return {"text": text, "confidence": 1.0, "bbox": []}


def load_docx(contents: bytes) -> tuple[list[dict], list[dict]]:
    """
    Extract text blocks and tables from a .docx file.

    Returns:
        (text_blocks, tables) where tables use the same cell dict format as
        table_detector (row, col, text, bbox=None).
    """
    import docx

    doc = docx.Document(io.BytesIO(contents))

    blocks = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            blocks.append(_block(text))

    tables = []
    for tbl in doc.tables:
        cells = []
        for r, row in enumerate(tbl.rows):
            for c, cell in enumerate(row.cells):
                text = cell.text.strip()
                if text:
                    cells.append({"row": r, "col": c, "text": text, "bbox": None})
        if cells:
            tables.append({"cells": cells, "bbox": None})
        # Also surface table content as plain text lines for full_text
        for row in tbl.rows:
            line = " | ".join(cell.text.strip() for cell in row.cells)
            if line.strip(" |"):
                blocks.append(_block(line))

    return blocks, tables


def load_xlsx(contents: bytes) -> tuple[list[dict], list[dict]]:
    """
    Extract text blocks and tables from an .xlsx file, one table per sheet.

    Returns:
        (text_blocks, tables) with the same formats as load_docx.
    """
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True, read_only=True)

    blocks = []
    tables = []
    for ws in wb.worksheets:
        cells = []
        for r, row in enumerate(ws.iter_rows(values_only=True)):
            values = ["" if v is None else str(v).strip() for v in row]
            if any(values):
                blocks.append(_block(f"[{ws.title}] " + "\t".join(values)))
                for c, text in enumerate(values):
                    if text:
                        cells.append({"row": r, "col": c, "text": text, "bbox": None})
        if cells:
            tables.append({"cells": cells, "bbox": None})

    wb.close()
    return blocks, tables
